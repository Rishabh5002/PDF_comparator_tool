"""Local Excel report writer."""
from __future__ import annotations
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font


def write_excel_report(payload: dict, output_path: str | Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Form PDF Comparison Report"])
    ws["A1"].font = Font(bold=True, size=16)
    ws.append(["Processing mode", "Offline / Local"])
    ws.append(["Original", payload["old_document"]["filename"]])
    ws.append(["Revision", payload["new_document"]["filename"]])
    ws.append([])
    ws.append(["Change type", "Count"])
    for key, value in payload["comparison"]["summary"].items():
        ws.append([key, value])

    diff = wb.create_sheet("Differences")
    headers = ["Type", "Question", "Original", "Revision", "Details", "Confidence"]
    diff.append(headers)
    for cell in diff[1]:
        cell.font = Font(bold=True)
    for d in payload["comparison"]["differences"]:
        diff.append([
            d.get("type"), d.get("question_number"), str(d.get("old_value", "")),
            str(d.get("new_value", "")), d.get("message", ""), d.get("confidence")
        ])

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            max_len = min(max(len(str(c.value or "")) for c in column) + 2, 60)
            sheet.column_dimensions[column[0].column_letter].width = max_len

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output
